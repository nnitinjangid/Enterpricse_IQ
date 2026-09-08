from pathlib import Path

import pandas as pd
import fitz
from docx import Document as DocxDocument


# =========================================================
# PDF
# =========================================================

def extract_text_from_pdf(file_path: str) -> str:
    """
    Extract text from all pages of a PDF.

    Page numbers are preserved so that later
    RAG responses can provide accurate citations.
    """

    text_parts = []

    pdf = fitz.open(file_path)

    try:
        for page_number, page in enumerate(pdf, start=1):

            page_text = page.get_text("text")

            if page_text.strip():
                text_parts.append(
                    f"[Page {page_number}]\n"
                    f"{page_text.strip()}"
                )

    finally:
        pdf.close()

    return "\n\n".join(text_parts)


# =========================================================
# DOCX
# =========================================================

def extract_text_from_docx(file_path: str) -> str:
    """
    Extract paragraphs and tables from DOCX.
    """

    document = DocxDocument(file_path)

    text_parts = []

    # -----------------------------------------------------
    # Paragraphs
    # -----------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # -----------------------------------------------------
    # Tables
    # -----------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            row_text = " | ".join(cells)

            if row_text.strip():
                text_parts.append(row_text)

    return "\n".join(text_parts)


# =========================================================
# CSV
# =========================================================

def extract_text_from_csv(file_path: str) -> str:
    """
    Convert CSV rows into readable text.
    """

    dataframe = pd.read_csv(file_path)

    return dataframe.to_string(
        index=False
    )


# =========================================================
# XLSX
# =========================================================

def extract_text_from_xlsx(file_path: str) -> str:
    """
    Extract data from all Excel sheets.
    """

    excel_file = pd.ExcelFile(file_path)

    text_parts = []

    for sheet_name in excel_file.sheet_names:

        dataframe = pd.read_excel(
            file_path,
            sheet_name=sheet_name,
        )

        sheet_text = (
            f"[Sheet: {sheet_name}]\n"
            f"{dataframe.to_string(index=False)}"
        )

        text_parts.append(sheet_text)

    return "\n\n".join(text_parts)


# =========================================================
# Universal Parser
# =========================================================

def extract_text(
    file_path: str,
    file_type: str,
) -> str:
    """
    Extract text based on document type.
    """

    file_type = (
        file_type
        .lower()
        .lstrip(".")
    )

    # -----------------------------------------------------
    # PDF
    # -----------------------------------------------------

    if file_type == "pdf":
        return extract_text_from_pdf(file_path)

    # -----------------------------------------------------
    # DOCX
    # -----------------------------------------------------

    if file_type == "docx":
        return extract_text_from_docx(file_path)

    # -----------------------------------------------------
    # CSV
    # -----------------------------------------------------

    if file_type == "csv":
        return extract_text_from_csv(file_path)

    # -----------------------------------------------------
    # XLSX
    # -----------------------------------------------------

    if file_type == "xlsx":
        return extract_text_from_xlsx(file_path)

    # -----------------------------------------------------
    # Unsupported file type
    # -----------------------------------------------------

    raise ValueError(
        f"Unsupported document type: {file_type}"
    )